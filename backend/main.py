

import os
import io
import json
import re
import pandas as pd

from docx import Document
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

from fastapi import FastAPI, HTTPException, UploadFile, File, Body, Response
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field

from backend.voyage_simulation import simulate_voyage
from backend.vessel_layout import create_vessel_layout
from backend.ai_service import ask_local_ai
from backend.models import CargoManifest
from backend.slots import VesselSlotPlan
from backend.placement import create_stowage_plan



app = FastAPI(
    title="AI-Driven Container Vessel Cargo Scheduling and Voyage Planning System",
    description="Local AI-assisted container stowage and voyage planning.",
    version="0.1.0",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=[],
    allow_origin_regex=r"^http://(localhost|127\.0\.0\.1):\d+$",
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)

class PlanningRequest(BaseModel):
    manifest: CargoManifest
    vessel: VesselSlotPlan | None = None
    use_ai: bool = False

class VoyageSimulationRequest(BaseModel):
    manifest: CargoManifest
    port_sequence: list[str] = Field(min_length=2)
    vessel: VesselSlotPlan | None = None

@app.get("/")
def home():
    return {
        "message": "Container Stowage API is running",
        "status": "ok",
    }


@app.get("/health")
def health_check():
    return {
        "status": "healthy",
        "service": "container-stowage-api",
        "ai_runtime": "Ollama",
    }


@app.post("/api/plan")
def generate_plan(request: PlanningRequest):
    if not request.manifest.containers:
        raise HTTPException(
            status_code=400,
            detail="Cargo manifest cannot be empty.",
        )

    try:
        vessel = request.vessel or create_vessel_layout()

        return create_stowage_plan(
            request.manifest,
            vessel,
            use_ai=request.use_ai,
        )        
    except ValueError as exc:
        raise HTTPException(
            status_code=400,
            detail=str(exc),
        ) from exc

from pydantic import BaseModel, Field


class AIQuestion(BaseModel):
    question: str = Field(min_length=1, max_length=2000)
    containers: list[dict] = Field(default_factory=list)
    plan: dict | None = None


def _build_ai_dataset_context(
    question: str,
    containers: list[dict],
    plan: dict | None,
) -> str | None:
    if not containers and not plan:
        return None

    question_lower = question.lower()

    # Find explicitly mentioned container IDs first.
    mentioned_ids = {
        str(item.get("container_id", "")).strip()
        for item in containers
        if str(item.get("container_id", "")).strip()
        and str(item.get("container_id", "")).lower() in question_lower
    }

    # If no container ID was mentioned, look for loading/discharge ports.
    known_ports = {
        str(item.get(key, "")).strip()
        for item in containers
        for key in ("loading_port", "discharge_port")
        if str(item.get(key, "")).strip()
    }

    mentioned_ports = {
        port for port in known_ports
        if port.lower() in question_lower
    }

    if mentioned_ids:
        relevant_containers = [
            item for item in containers
            if str(item.get("container_id", "")).strip() in mentioned_ids
        ]
    elif mentioned_ports:
        relevant_containers = [
            item for item in containers
            if str(item.get("loading_port", "")).strip() in mentioned_ports
            or str(item.get("discharge_port", "")).strip() in mentioned_ports
        ]
    elif len(containers) <= 100:
        relevant_containers = containers
    else:
        relevant_containers = containers[:100]

    # Keep plan assignments aligned with the relevant containers.
    relevant_ids = {
        str(item.get("container_id", "")).strip()
        for item in relevant_containers
    }

    plan_context = None
    if plan:
        assignments = plan.get("assignments", [])
        relevant_assignments = [
            item for item in assignments
            if not relevant_ids
            or str(item.get("container_id", "")).strip() in relevant_ids
        ]

        plan_context = {
            "vessel_id": plan.get("vessel_id"),
            "assigned_count": plan.get("assigned_count"),
            "unassigned_count": plan.get("unassigned_count"),
            "estimated_rehandles": plan.get("estimated_rehandles"),
            "assignments": relevant_assignments[:100],
        }

    destination_counts = {}
    for item in containers:
        port = str(item.get("discharge_port", "")).strip()
        if port:
            destination_counts[port] = destination_counts.get(port, 0) + 1

    context = {
        "dataset_summary": {
            "total_containers": len(containers),
            "discharge_port_counts": destination_counts,
            "context_rows_returned": len(relevant_containers),
            "context_truncated": (
                len(containers) > len(relevant_containers)
                and not mentioned_ids
                and not mentioned_ports
            ),
        },
        "containers": relevant_containers,
        "latest_stowage_plan": plan_context,
    }

    return json.dumps(context, indent=2, default=str)


def _fast_dataset_answer(
    question: str,
    containers: list[dict],
    plan: dict | None,
) -> str | None:
    """Answer common factual dataset questions without calling the LLM."""

    if not containers:
        return None

    q = question.lower().strip()

    container_by_id = {
        str(item.get("container_id", "")).strip().upper(): item
        for item in containers
        if str(item.get("container_id", "")).strip()
    }

    assignment_by_id = {}
    if plan:
        assignment_by_id = {
            str(item.get("container_id", "")).strip().upper(): item
            for item in plan.get("assignments", [])
            if str(item.get("container_id", "")).strip()
        }

    # Dataset-level container count.
    if (
        ("how many" in q or "number of" in q or "count" in q)
        and "container" in q
        and not any(port.lower() in q for port in {
            str(item.get("discharge_port", "")).strip()
            for item in containers
            if str(item.get("discharge_port", "")).strip()
        })
    ):
        return f"There are {len(containers)} containers in the uploaded dataset."

    # Counts by discharge port.
    discharge_ports = {
        str(item.get("discharge_port", "")).strip()
        for item in containers
        if str(item.get("discharge_port", "")).strip()
    }

    for port in discharge_ports:
        if port.lower() in q and (
            "how many" in q
            or "number of" in q
            or "count" in q
            or "going" in q
            or "unload" in q
            or "discharge" in q
        ):
            count = sum(
                1 for item in containers
                if str(item.get("discharge_port", "")).strip() == port
            )
            return f"{count} containers are scheduled to discharge at {port}."

    # Find a mentioned container ID such as CONT-0007.
    match = re.search(r"\b[A-Z]+[-_]?\d+\b", question.upper())
    if not match:
        return None

    container_id = match.group(0).replace("_", "-")
    container = container_by_id.get(container_id)

    if not container:
        return (
            f"{container_id} was not found in the currently uploaded dataset."
        )

    assignment = assignment_by_id.get(container_id)

    weight = container.get("weight_kg")
    size = container.get("size_ft")
    loading_port = container.get("loading_port")
    discharge_port = container.get("discharge_port")
    destination_order = container.get("destination_order", 1)
    refrigerated = bool(container.get("refrigerated", False))
    hazardous = bool(container.get("hazardous", False))

    if "weight" in q or "weigh" in q:
        return f"{container_id} weighs {weight} kg."

    if (
        "where" in q
        and ("unload" in q or "discharge" in q or "going" in q)
    ) or "discharge port" in q:
        return f"{container_id} will discharge at {discharge_port}."

    if "loading port" in q or ("where" in q and "load" in q):
        return f"{container_id} is loaded at {loading_port}."

    if "size" in q:
        return f"{container_id} is a {size}-ft container."

    if "destination order" in q or "priority" in q:
        return (
            f"{container_id} has destination order {destination_order}."
        )

    if "refrigerated" in q or "reefer" in q:
        return (
            f"{container_id} is "
            f"{'refrigerated' if refrigerated else 'not refrigerated'}."
        )

    if "hazardous" in q or "dangerous" in q:
        return (
            f"{container_id} is "
            f"{'hazardous cargo' if hazardous else 'not marked hazardous'}."
        )

    if "slot" in q or "position" in q or "assigned" in q:
        if not assignment:
            return (
                f"No generated stowage assignment is currently available "
                f"for {container_id}. Generate a stowage plan first."
            )

        if "why" in q:
            reason = assignment.get(
                "assignment_reason",
                "No assignment reason is available.",
            )
            details = assignment.get("reason_details", [])
            if details:
                return reason + " " + " ".join(str(item) for item in details)
            return reason

        return (
            f"{container_id} is assigned to position "
            f"{assignment.get('position_code', 'N/A')} "
            f"(slot {assignment.get('slot_id', 'N/A')})."
        )

    # General request about one container.
    return (
        f"{container_id}: {size}-ft, {weight} kg, loaded at {loading_port}, "
        f"discharged at {discharge_port}, destination order {destination_order}, "
        f"refrigerated: {'yes' if refrigerated else 'no'}, "
        f"hazardous: {'yes' if hazardous else 'no'}."
    )


@app.post("/api/ai/ask")
def ask_ai(request: AIQuestion):
    try:
        fast_answer = _fast_dataset_answer(
            request.question,
            request.containers,
            request.plan,
        )

        if fast_answer is not None:
            return {
                "model": "dataset-engine",
                "answer": fast_answer,
                "dataset_context_used": True,
                "uploaded_container_count": len(request.containers),
                "response_mode": "instant_dataset_lookup",
            }

        dataset_context = _build_ai_dataset_context(
            request.question,
            request.containers,
            request.plan,
        )

        answer = ask_local_ai(
            request.question,
            dataset_context=dataset_context,
        )

        return {
            "model": "qwen2.5:7b",
            "answer": answer,
            "dataset_context_used": bool(dataset_context),
            "uploaded_container_count": len(request.containers),
            "response_mode": "local_ai",
        }
    except Exception as exc:
        if os.getenv("VERCEL"):
            return {
                "model": "dataset-engine",
                "answer": (
                    "The cloud deployment can answer uploaded dataset and "
                    "generated-plan questions, but general Qwen/Ollama reasoning "
                    "is only available when the project is running locally with Ollama."
                ),
                "dataset_context_used": bool(request.containers or request.plan),
                "uploaded_container_count": len(request.containers),
                "response_mode": "cloud_dataset_only",
            }

        raise HTTPException(
            status_code=503,
            detail=f"Local AI service unavailable: {exc}",
        ) from exc


@app.post("/api/manifest/upload")
async def upload_manifest(file: UploadFile = File(...)):
    filename = (file.filename or "").lower()

    try:
        content = await file.read()

        if filename.endswith(".csv"):
            df = pd.read_csv(io.BytesIO(content))
        elif filename.endswith(".xlsx"):
            df = pd.read_excel(io.BytesIO(content), engine="openpyxl")
        else:
            raise HTTPException(
                status_code=400,
                detail="Upload a CSV (.csv) or Excel (.xlsx) file.",
            )

        required_columns = [
            "container_id",
            "size_ft",
            "weight_kg",
            "loading_port",
            "discharge_port",
        ]

        missing = [
            column for column in required_columns
            if column not in df.columns
        ]

        if missing:
            raise HTTPException(
                status_code=400,
                detail=f"Missing required columns: {', '.join(missing)}",
            )

        df = df.where(pd.notna(df), None)       

        return {
            "filename": file.filename,
            "total_rows": len(df),
            "columns": list(df.columns),
            "preview": df.head(10).to_dict(orient="records"),
            "containers": df.to_dict(orient="records"),
            "message": "File uploaded and required columns checked.",
        }


    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=f"Could not read the uploaded file: {exc}",
        )
    finally:
        await file.close()

@app.post("/api/voyage/simulate")
def simulate_voyage_endpoint(request: VoyageSimulationRequest):
    try:
        vessel = request.vessel or create_vessel_layout()

        return simulate_voyage(
            manifest=request.manifest,
            vessel=vessel,
            port_sequence=request.port_sequence,
        )
    except ValueError as exc:
        raise HTTPException(
            status_code=400,
            detail=str(exc),
        ) from exc




def _plan_assignments_dataframe(plan: dict) -> pd.DataFrame:
    assignments = plan.get("assignments", [])

    rows = []
    for item in assignments:
        rows.append({
            "container_id": item.get("container_id", ""),
            "slot_id": item.get("slot_id", ""),
            "position_code": item.get("position_code", ""),
            "bay": item.get("bay", ""),
            "row": item.get("row", ""),
            "tier": item.get("tier", ""),
            "size_ft": item.get("size_ft", ""),
            "selection_mode": item.get("selection_mode", ""),
            "assignment_reason": item.get("assignment_reason", ""),
        })

    return pd.DataFrame(rows)


@app.post("/api/plan/export/{file_format}")
def export_plan(
    file_format: str,
    plan: dict = Body(...),
):
    """Export the current generated stowage plan without storing it."""

    file_format = file_format.lower()
    supported = {"csv", "xlsx", "json", "docx", "pdf"}

    if file_format not in supported:
        raise HTTPException(
            status_code=400,
            detail="Supported formats: csv, xlsx, json, docx, pdf.",
        )

    df = _plan_assignments_dataframe(plan)
    base_name = "stowage_plan"

    if file_format == "csv":
        content = df.to_csv(index=False).encode("utf-8")
        media_type = "text/csv"
        filename = f"{base_name}.csv"

    elif file_format == "xlsx":
        buffer = io.BytesIO()
        with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Assignments")

            summary = pd.DataFrame([
                ["Vessel ID", plan.get("vessel_id", "")],
                ["Assigned containers", plan.get("assigned_count", 0)],
                ["Unassigned containers", plan.get("unassigned_count", 0)],
                ["Estimated rehandles", plan.get("estimated_rehandles", 0)],
                ["AI proposals", plan.get("ai_proposed_count", 0)],
                ["AI accepted", plan.get("ai_accepted_count", 0)],
                ["AI fallbacks", plan.get("ai_fallback_count", 0)],
            ], columns=["Metric", "Value"])
            summary.to_excel(writer, index=False, sheet_name="Summary")

        content = buffer.getvalue()
        media_type = (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        filename = f"{base_name}.xlsx"

    elif file_format == "json":
        content = json.dumps(plan, indent=2, default=str).encode("utf-8")
        media_type = "application/json"
        filename = f"{base_name}.json"

    elif file_format == "docx":
        document = Document()
        document.add_heading(
            "AI-Driven Container Vessel Cargo Scheduling and Voyage Planning System",
            level=1,
        )
        document.add_heading("Saved Stowage Plan", level=2)

        document.add_paragraph(
            f"Vessel ID: {plan.get('vessel_id', '')}"
        )
        document.add_paragraph(
            f"Assigned containers: {plan.get('assigned_count', 0)}"
        )
        document.add_paragraph(
            f"Unassigned containers: {plan.get('unassigned_count', 0)}"
        )
        document.add_paragraph(
            f"Estimated rehandles: {plan.get('estimated_rehandles', 0)}"
        )

        if not df.empty:
            table = document.add_table(rows=1, cols=len(df.columns))
            table.style = "Table Grid"

            for index, column in enumerate(df.columns):
                table.rows[0].cells[index].text = str(column)

            for _, row in df.fillna("").iterrows():
                cells = table.add_row().cells
                for index, value in enumerate(row):
                    cells[index].text = str(value)

        buffer = io.BytesIO()
        document.save(buffer)
        content = buffer.getvalue()
        media_type = (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        filename = f"{base_name}.docx"

    else:
        buffer = io.BytesIO()
        pdf = SimpleDocTemplate(
            buffer,
            pagesize=landscape(A4),
            rightMargin=10 * mm,
            leftMargin=10 * mm,
            topMargin=10 * mm,
            bottomMargin=10 * mm,
        )
        styles = getSampleStyleSheet()
        story = [
            Paragraph(
                "AI-Driven Container Vessel Cargo Scheduling and Voyage Planning System",
                styles["Title"],
            ),
            Spacer(1, 6),
            Paragraph("Saved Stowage Plan", styles["Heading2"]),
            Spacer(1, 6),
            Paragraph(
                (
                    f"Vessel: {plan.get('vessel_id', '')} &nbsp;&nbsp; "
                    f"Assigned: {plan.get('assigned_count', 0)} &nbsp;&nbsp; "
                    f"Unassigned: {plan.get('unassigned_count', 0)} &nbsp;&nbsp; "
                    f"Estimated rehandles: {plan.get('estimated_rehandles', 0)}"
                ),
                styles["BodyText"],
            ),
            Spacer(1, 10),
        ]

        if not df.empty:
            pdf_columns = [
                "container_id",
                "position_code",
                "size_ft",
                "selection_mode",
                "assignment_reason",
            ]
            pdf_df = df[pdf_columns].fillna("")
            table_data = [pdf_columns] + pdf_df.astype(str).values.tolist()
            table = Table(
                table_data,
                repeatRows=1,
                colWidths=[32 * mm, 28 * mm, 18 * mm, 32 * mm, 125 * mm],
            )
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#111827")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [
                    colors.white,
                    colors.HexColor("#f8fafc"),
                ]),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.append(table)
        else:
            story.append(
                Paragraph("No assignments available in this plan.", styles["BodyText"])
            )

        pdf.build(story)
        content = buffer.getvalue()
        media_type = "application/pdf"
        filename = f"{base_name}.pdf"

    return Response(
        content=content,
        media_type=media_type,
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Cache-Control": "no-store",
        },
    )


@app.get("/api/vessel/layout")
def get_vessel_layout():
    vessel = create_vessel_layout()
    data = vessel.model_dump()

    for slot_data, slot in zip(data["slots"], vessel.slots):
        slot_data["position_code"] = slot.position_code

    return data

